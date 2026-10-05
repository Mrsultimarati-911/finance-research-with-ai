#!/usr/bin/env python3
"""Scoped, read-only local evidence lookup. Apache-2.0.

Original adapter, 2026-10-05. PDF page signals adapt the project-local
paper-deepread/scripts/prepare_pdf.py; no dependency on that Skill is required.
No macros, formulas, links, source instructions or executable files are run.
"""
from __future__ import annotations

import argparse
import codecs
from collections import deque
import csv
from datetime import datetime, timezone
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import stat
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".py", ".r", ".js", ".ts", ".jsx", ".tsx", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".log", ".sql", ".html", ".css", ".xml", ".sh", ".ps1", ".bat", ".tex", ".bib", ".ris", ".c", ".cpp", ".h", ".ipynb"}
SUPPORTED = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx", ".xlsm", ".csv"}
EXCLUDED_DIRECTORIES = {".git", ".venv", "venv", "node_modules", "__pycache__", ".pytest_cache", ".build"}


def now():
    return datetime.now(timezone.utc).isoformat()


def scoped_path(root, path=None):
    """Resolve before opening; reject sibling-prefix paths and escaping links."""
    root = Path(root).expanduser().resolve(strict=True)
    if not root.is_dir():
        raise ValueError("root must be an existing directory")
    candidate = root if path is None else Path(path).expanduser()
    if not candidate.is_absolute():
        candidate = root / candidate
    candidate = candidate.resolve(strict=True)
    if not candidate.is_relative_to(root):
        raise ValueError("path resolves outside the allowed root")
    return root, candidate


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def ensure_output_distinct(output, input_paths):
    """Reject source aliases before writing; independent result files may recur."""
    target = Path(output).expanduser().resolve()
    for value in input_paths:
        if value is None:
            continue
        source = Path(value).expanduser().resolve()
        same_file = target == source
        if not same_file and target.is_file() and source.is_file():
            same_file = target.samefile(source)
        if same_file:
            raise ValueError("output must not overwrite an input source file")


def base_result(operation, root):
    return {"schema_version": 2, "operation": operation, "root": str(root),
            "status": "completed", "retrieved_at": now(), "evidence": [],
            "files": [], "unprocessed": [], "coverage": {},
            "limitations": ["Extraction and anchor checks do not certify semantic support.",
                            "Publication dates are not inferred from file modification times."]}


def inventory(root, include_excluded=False):
    root, _ = scoped_path(root)
    result = base_result("inventory", root)
    records = []
    errors = []

    def walk_error(error):
        errors.append({"path": error.filename, "status": "permission_or_io_error", "reason": str(error)})

    for directory, folders, filenames in os.walk(root, followlinks=False, onerror=walk_error):
        for name in list(folders):
            path = Path(directory) / name
            try:
                resolved = path.resolve(strict=True)
                reparse_directory = bool(getattr(path.lstat(), "st_file_attributes", 0) & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0))
                reason = "excluded_symlink_or_junction" if path.is_symlink() or reparse_directory or not resolved.is_relative_to(root) or (hasattr(path, "is_junction") and path.is_junction()) else "excluded_directory" if name in EXCLUDED_DIRECTORIES and not include_excluded else None
            except OSError as exc:
                reason = "resolution_error"
                errors.append({"path": str(path), "status": reason, "reason": str(exc)})
            if reason:
                folders.remove(name)
                records.append({"path": str(path), "relative_path": str(path.relative_to(root)), "status": reason, "supported": False, "kind": "directory"})
        for name in sorted(filenames):
            path = Path(directory) / name
            try:
                _, resolved = scoped_path(root, path)
                file_stat = resolved.stat()
                supported = resolved.suffix.lower() in SUPPORTED or resolved.name.lower() in {"readme", "license", "makefile", "dockerfile", ".env"}
                records.append({"path": str(resolved), "relative_path": str(path.relative_to(root)),
                                "extension": resolved.suffix.lower(), "bytes": file_stat.st_size,
                                "modified_at": datetime.fromtimestamp(file_stat.st_mtime, timezone.utc).isoformat(),
                                "status": "discovered" if supported else "unsupported_format", "supported": supported, "kind": "file"})
            except (OSError, ValueError) as exc:
                records.append({"path": str(path), "relative_path": str(path.relative_to(root)), "status": "outside_scope_or_io_error", "supported": False, "kind": "file", "reason": str(exc)})
    result["files"] = sorted(records, key=lambda item: item["relative_path"].casefold())
    result["unprocessed"] = errors + [record for record in records if record.get("kind") == "file" and not record["supported"]]
    result["coverage"] = {"eligible_files": sum(record["supported"] for record in records), "inventory_entries": len(records), "excluded_directories": sorted(EXCLUDED_DIRECTORIES) if not include_excluded else [], "symlink_directories_followed": False}
    if errors:
        result["status"] = "partial"
    return result


def detect_encoding(path, override=None):
    """Validate incrementally; do not load large CSV/text files into memory."""
    with Path(path).open("rb") as source:
        prefix = source.read(4096)
    if override:
        candidates = [override]
    elif prefix.startswith(codecs.BOM_UTF16_LE) or prefix.startswith(codecs.BOM_UTF16_BE):
        candidates = ["utf-16"]
    else:
        if b"\x00" in prefix:
            raise ValueError("binary or unsupported encoding; provide an explicit encoding for text")
        candidates = ["utf-8-sig", "gb18030"]
    for encoding in candidates:
        try:
            decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
            with Path(path).open("rb") as source:
                for block in iter(lambda: source.read(1024 * 1024), b""):
                    decoder.decode(block)
                decoder.decode(b"", final=True)
            return encoding
        except (UnicodeError, LookupError):
            continue
    raise ValueError("cannot decode file without replacement; use --encoding with a known encoding")


def pdf_signal(text):
    visible = sum(not char.isspace() for char in text)
    signal = "empty" if visible == 0 else "sparse" if visible < 30 else "text_present"
    if text.count("\ufffd") / max(visible, 1) > .05:
        signal = "encoding_suspect"
    return {"nonspace_characters": visible, "signal": signal, "needs_review": signal != "text_present"}


def units(path, details, encoding=None, csv_header=True):
    extension = path.suffix.lower()
    if extension == ".pdf":
        engine = "pdfplumber" if importlib.util.find_spec("pdfplumber") else "pypdf"
        if not importlib.util.find_spec(engine):
            raise RuntimeError("missing_optional_dependency: pdfplumber or pypdf")
        details.update({"engine": engine, "page_signals": [], "printed_page_numbers": "not_verified", "visual_layout": "not_verified"})
        if engine == "pdfplumber":
            import pdfplumber
            with pdfplumber.open(path) as document:
                details["page_count"] = len(document.pages)
                for index, page in enumerate(document.pages, 1):
                    text = page.extract_text(layout=False) or ""
                    signal = {"physical_page": index, **pdf_signal(text)}
                    details["page_signals"].append(signal)
                    yield {"type": "pdf", "physical_page": index, "extracted_line_start": 1, "extracted_line_end": len(text.splitlines())}, text, signal
        else:
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("encrypted_pdf: password or authorized host reader required")
            details["page_count"] = len(reader.pages)
            for index, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                signal = {"physical_page": index, **pdf_signal(text)}
                details["page_signals"].append(signal)
                yield {"type": "pdf", "physical_page": index, "extracted_line_start": 1, "extracted_line_end": len(text.splitlines())}, text, signal
    elif extension == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        paragraph_index = table_index = 0
        section = []
        details.update({"engine": "python-docx", "read_scope": "body paragraphs and top-level table cells", "not_read": ["headers", "footers", "comments", "tracked revisions", "embedded objects"], "page_numbers": "not_rendered"})
        for element in document.element.body.iterchildren():
            if element.tag.endswith("}p"):
                paragraph_index += 1
                paragraph = Paragraph(element, document)
                style = paragraph.style.name if paragraph.style else ""
                match = re.match(r"(?:Heading|标题)\s*(\d+)", style, re.I)
                if match:
                    level = int(match[1])
                    section = section[:level - 1]
                    section.append(paragraph.text)
                yield {"type": "docx_paragraph", "paragraph_index": paragraph_index, "section": list(section)}, paragraph.text, {"style": style}
            elif element.tag.endswith("}tbl"):
                table_index += 1
                table = Table(element, document)
                for row_index, row in enumerate(table.rows, 1):
                    for column_index, cell in enumerate(row.cells, 1):
                        text = cell.text
                        nested = []
                        for nested_table in cell.tables:
                            for nested_row in nested_table.rows:
                                nested.append(" | ".join(c.text for c in nested_row.cells))
                        if nested:
                            text += "\n" + "\n".join(nested)
                        yield {"type": "docx_table_cell", "table_index": table_index, "row": row_index, "column": column_index, "section": list(section)}, text, {"grid_coordinate": True, "nested_tables_included_as_text": bool(nested)}
    elif extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        # XML metadata and read-only worksheets avoid loading every cell at once.
        structure = workbook_structure(path)
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        details.update({"engine": "openpyxl", "formulas_evaluated": False, "external_links_followed": False, "macros_executed": False, "worksheets": []})
        try:
            for worksheet in formulas.worksheets:
                cache_sheet = cached[worksheet.title]
                sheet_metadata = structure.get(worksheet.title, {"merged_ranges": [], "hidden_rows": [], "hidden_column_ranges": []})
                ranges = sheet_metadata["merged_ranges"]
                hidden_rows = set(sheet_metadata["hidden_rows"])
                hidden_columns = sheet_metadata["hidden_column_ranges"]
                details["worksheets"].append({"name": worksheet.title, "state": worksheet.sheet_state, **sheet_metadata})
                bounds = [(name, range_boundaries(name)) for name in ranges]
                first_row = None
                for cells, cache_cells in zip(worksheet.iter_rows(), cache_sheet.iter_rows()):
                    context = {cell.coordinate: cell.value for cell in cells if cell.value is not None}
                    if first_row is None:
                        first_row = dict(context)
                    for cell, cache_cell in zip(cells, cache_cells):
                        if cell.value is None:
                            continue
                        formula = cell.value if cell.data_type == "f" else None
                        value = cache_cell.value if formula else cell.value
                        region = next((name for name, (c1, r1, c2, r2) in bounds if c1 <= cell.column <= c2 and r1 <= cell.row <= r2), None)
                        metadata = {"formula": formula, "cached_value": value if formula else None, "value": value,
                                    "cache_status": "missing" if formula and value is None else "present_unverified_freshness" if formula else "not_applicable",
                                    "number_format": cell.number_format, "sheet_state": worksheet.sheet_state, "merged_range": region,
                                    "row_hidden": cell.row in hidden_rows, "column_hidden": any(c1 <= cell.column <= c2 for c1, c2 in hidden_columns),
                                    "row_context": context, "first_row_context": first_row, "unit": "read_from_context_not_inferred"}
                        text = str(cell.value)
                        if formula:
                            text += "\nCached value: " + ("UNKNOWN" if value is None else str(value))
                        yield {"type": "xlsx_cell", "sheet": worksheet.title, "cell": cell.coordinate}, text, metadata
        finally:
            formulas.close()
            cached.close()
    elif extension == ".csv":
        selected = detect_encoding(path, encoding)
        details.update({"engine": "csv", "encoding": selected, "encoding_detection": "explicit" if encoding else "validated_heuristic", "streaming": True})
        with path.open("r", encoding=selected, newline="") as source:
            sample = source.read(8192)
            source.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(source, dialect)
            first_row = next(reader, [])
            headers = first_row if csv_header else [f"column_{index}" for index in range(1, len(first_row) + 1)]
            if not csv_header:
                source.seek(0)
                reader = csv.reader(source, dialect)
            details.update({"columns": headers, "duplicate_columns": sorted({name for name in headers if headers.count(name) > 1}), "header_physical_line_end": reader.line_num, "delimiter": dialect.delimiter, "header_present": csv_header, "header_policy": "first record as header" if csv_header else "explicit headerless input"})
            prior = reader.line_num
            for index, row in enumerate(reader, 1):
                locator = {"type": "csv_record", "record_number": index, "physical_line_start": prior + 1, "physical_line_end": reader.line_num, "columns": headers}
                prior = reader.line_num
                text = json.dumps({"columns": headers, "values": row}, ensure_ascii=False)
                yield locator, text, {"values": row, "column_count_matches": len(row) == len(headers), "encoding": selected, "header_present": csv_header}
    else:
        selected = detect_encoding(path, encoding)
        details.update({"engine": "text", "encoding": selected, "encoding_detection": "explicit" if encoding else "validated_heuristic", "streaming": True})
        with path.open("r", encoding=selected, newline="") as source:
            previous, upcoming = deque(maxlen=2), deque()
            iterator = iter(enumerate(source, 1))
            for item in iterator:
                upcoming.append((item[0], item[1].rstrip("\r\n")))
                if len(upcoming) >= 3:
                    index, line = upcoming.popleft()
                    context = [{"line": number, "text": text} for number, text in list(previous) + [(index, line)] + list(upcoming)]
                    yield {"type": "text", "line_start": index, "line_end": index}, line, {"encoding": selected, "context": context}
                    previous.append((index, line))
            while upcoming:
                index, line = upcoming.popleft()
                context = [{"line": number, "text": text} for number, text in list(previous) + [(index, line)] + list(upcoming)]
                yield {"type": "text", "line_start": index, "line_end": index}, line, {"encoding": selected, "context": context}
                previous.append((index, line))


def workbook_structure(path):
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    result = {}
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
        for sheet in workbook.findall("s:sheets/s:sheet", ns):
            target = targets[sheet.attrib["{" + ns["r"] + "}id"]]
            member = target.lstrip("/") if target.startswith("/") else "xl/" + target
            merges, hidden_rows, hidden_columns = [], [], []
            with archive.open(member) as source:
                for _, element in ET.iterparse(source, events=("end",)):
                    if element.tag.endswith("}mergeCell"):
                        merges.append(element.attrib["ref"])
                    elif element.tag.endswith("}row") and element.attrib.get("hidden") in {"1", "true"}:
                        hidden_rows.append(int(element.attrib["r"]))
                    elif element.tag.endswith("}col") and element.attrib.get("hidden") in {"1", "true"}:
                        hidden_columns.append([int(element.attrib["min"]), int(element.attrib["max"])])
                    element.clear()
            result[sheet.attrib["name"]] = {"merged_ranges": merges, "hidden_rows": hidden_rows, "hidden_column_ranges": hidden_columns}
    return result


def evidence_record(path, sha, locator, excerpt, metadata, retrieved_at):
    key = str(path) + sha + json.dumps(locator, sort_keys=True, ensure_ascii=False)
    return {"evidence_id": "L-" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:16], "source_type": "local_file",
            "source_name": path.name, "source_path": str(path), "sha256": sha,
            "publication_date": None, "retrieved_at": retrieved_at, "locator": locator, "excerpt": excerpt,
            "read_range": locator, "verification_status": "needs_review" if metadata.get("needs_review") else "extracted", "metadata": metadata}


def read(root, path, start=0, count=None, encoding=None, csv_header=True):
    root, path = scoped_path(root, path)
    if not path.is_file():
        raise ValueError("path must be a file")
    if start < 0 or (count is not None and count < 1):
        raise ValueError("start must be nonnegative and count positive")
    result = base_result("read", root)
    details = {"path": str(path), "status": "reading", "modified_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()}
    result["files"].append(details)
    if path.suffix.lower() not in SUPPORTED and path.name.lower() not in {"readme", "license", "makefile", "dockerfile", ".env"}:
        details.update({"status": "unsupported_format", "reason": "No reader is registered for this file type"})
        result["status"] = "failed"
        result["unprocessed"].append(dict(details))
        return result
    sha = file_hash(path)
    details["sha256"] = sha
    total = 0
    try:
        for index, (locator, text, metadata) in enumerate(units(path, details, encoding, csv_header)):
            total = index + 1
            if index < start or (count is not None and index >= start + count):
                continue
            result["evidence"].append(evidence_record(path, sha, locator, text, metadata, result["retrieved_at"]))
        if file_hash(path) != sha:
            raise RuntimeError("source_changed_during_read; repeat the read before using these excerpts")
        details.update({"status": "read", "units_total": total, "units_returned": len(result["evidence"])})
        pending = []
        if start:
            pending.append({"path": str(path), "status": "not_returned", "unit_start": 0, "unit_end_exclusive": min(start, total), "reason": "requested_unit_range"})
        end = min(total, start + count) if count is not None else total
        if end < total:
            pending.append({"path": str(path), "status": "not_returned", "unit_start": end, "unit_end_exclusive": total, "reason": "requested_unit_range"})
        review_pages = [page["physical_page"] for page in details.get("page_signals", []) if page["needs_review"]]
        review_tasks = [{"path": str(path), "status": "requires_visual_review", "physical_page": page, "reason": "empty_sparse_or_encoding_suspect_pdf_text; OCR or visual reading is not performed by this helper"} for page in review_pages]
        result["unprocessed"] = pending + review_tasks
        result["status"] = "partial" if result["unprocessed"] else "completed"
        result["coverage"] = {"units_total": total, "requested_unit_start": start, "returned_unit_start": min(start, total), "returned_unit_end_exclusive": end, "all_units_scanned": True, "all_units_returned": not pending, "unverified_pages": review_pages, "visual_layout_verified": False if path.suffix.lower() == ".pdf" else "not_applicable"}
    except Exception as exc:
        details.update({"status": "read_error", "reason": f"{type(exc).__name__}: {exc}", "units_scanned_before_failure": total})
        result["status"] = "partial" if result["evidence"] else "failed"
        result["unprocessed"].append(dict(details))
        result["coverage"] = {"all_units_scanned": False, "units_scanned_before_failure": total}
    return result


def rg_candidates(paths, queries):
    """Use rg as the first text candidate pass; never invoke a command shell."""
    executable = shutil.which("rg")
    if not executable:
        return None, {"engine": "python_fallback", "reason": "rg_unavailable"}
    found = set()
    # Batch arguments to stay below Windows command-line length limits.
    for start in range(0, len(paths), 40):
        command = [executable, "--json", "--fixed-strings", "--ignore-case", "--text"]
        for query in queries:
            command += ["-e", query]
        command += ["--"] + [str(path) for path in paths[start:start + 40]]
        run = subprocess.run(command, capture_output=True, encoding="utf-8", errors="replace", shell=False)
        if run.returncode not in (0, 1):
            return None, {"engine": "python_fallback", "reason": "rg_error", "detail": run.stderr.strip()}
        for line in run.stdout.splitlines():
            event = json.loads(line)
            if event.get("type") == "match":
                data = event["data"]["path"]
                if "text" in data:
                    found.add(str(Path(data["text"]).resolve()))
    return found, {"engine": "rg_then_format_reader", "rg_available": True}


def search(root, queries, include_excluded=False, encoding=None, csv_header=True):
    queries = list(dict.fromkeys(str(query).strip() for query in queries if str(query).strip()))
    if not queries:
        raise ValueError("at least one nonempty query is required")
    discovered = inventory(root, include_excluded)
    result = base_result("search", discovered["root"])
    result["queries"] = queries
    result["limitations"].append("Text matches are within original lines; phrases spanning line breaks require separate terms or a targeted read. Context preserves up to two original lines on either side.")
    result["unprocessed"] = list(discovered["unprocessed"])
    eligible = [item for item in discovered["files"] if item["supported"]]
    text_paths = [Path(item["path"]) for item in eligible if Path(item["path"]).suffix.lower() not in {".csv", ".pdf", ".docx", ".xlsx", ".xlsm"}]
    candidate_paths, candidate_details = rg_candidates(text_paths, queries) if text_paths else (set(), {"engine": "format_readers"})
    folded = [query.casefold() for query in queries]
    eligible.sort(key=lambda item: (candidate_paths is not None and item["path"] not in candidate_paths, item["relative_path"].casefold()))
    for item in eligible:
        path = Path(item["path"])
        file_details = {**item, "status": "reading", "rg_candidate": None if candidate_paths is None else str(path) in candidate_paths}
        matches = []
        unit_count = 0
        try:
            # Re-check scope before every file, and keep only matching units.
            # rg candidates are prioritized; a streaming fallback scans every
            # eligible file for non-UTF8 and Unicode case-folding completeness.
            _, path = scoped_path(result["root"], path)
            sha = file_hash(path)
            file_details["sha256"] = sha
            for locator, text, metadata in units(path, file_details, encoding, csv_header):
                unit_count += 1
                hit_queries = [query for query, token in zip(queries, folded) if token in text.casefold()]
                if hit_queries:
                    evidence = evidence_record(path, sha, locator, text, metadata, result["retrieved_at"])
                    evidence["matched_queries"] = hit_queries
                    matches.append(evidence)
            if file_hash(path) != sha:
                raise RuntimeError("source_changed_during_read; repeat the search before using these excerpts")
            file_details.update({"status": "read", "units_scanned": unit_count, "all_units_scanned": True})
            for page in file_details.get("page_signals", []):
                if page["needs_review"]:
                    result["unprocessed"].append({"path": str(path), "status": "requires_visual_review", "physical_page": page["physical_page"], "reason": "page text is empty, sparse or encoding-suspect; not searched visually"})
        except Exception as exc:
            file_details.update({"status": "read_error", "reason": f"{type(exc).__name__}: {exc}", "units_scanned_before_failure": unit_count, "all_units_scanned": False})
            result["unprocessed"].append(dict(file_details))
        file_details["match_count"] = len(matches)
        if file_details["status"] == "read":
            file_details["status"] = "searched_matches" if matches else "searched_no_match"
        result["files"].append(file_details)
        result["evidence"].extend(matches)
    result["coverage"] = {**discovered["coverage"], **candidate_details, "files_attempted": len(eligible), "evidence_matches": len(result["evidence"]), "fixed_result_cap": None,
                          "queries_executed": queries, "source_scope": "local_root_only", "internet_calls": 0,
                          "text_fallback": "all eligible text files scanned incrementally to preserve encoding and Unicode matches"}
    if result["unprocessed"]:
        result["status"] = "partial" if result["evidence"] or any(item.get("status") in {"searched_matches", "searched_no_match"} for item in result["files"]) else "failed"
    elif not result["evidence"]:
        result["status"] = "no_results"
    return result


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("operation", choices=["inventory", "search", "read"])
    parser.add_argument("--root", required=True)
    parser.add_argument("--path")
    parser.add_argument("--query", action="append", default=[])
    parser.add_argument("--start", type=int, default=0)
    parser.add_argument("--count", type=int)
    parser.add_argument("--encoding")
    parser.add_argument("--csv-no-header", action="store_true", help="Treat every CSV record as data; columns become column_1, column_2, etc.")
    parser.add_argument("--include-excluded", action="store_true")
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    try:
        result = inventory(args.root, args.include_excluded) if args.operation == "inventory" else search(args.root, args.query, args.include_excluded, args.encoding, not args.csv_no_header) if args.operation == "search" else read(args.root, args.path, args.start, args.count, args.encoding, not args.csv_no_header) if args.path else (_ for _ in ()).throw(ValueError("read requires --path"))
        content = json.dumps(result, ensure_ascii=False, indent=2, default=str) + "\n"
        if args.out:
            ensure_output_distinct(args.out, [item.get("path") for item in result["files"]] + [item.get("path") for item in result["unprocessed"]])
            args.out.parent.mkdir(parents=True, exist_ok=True)
            args.out.write_text(content, encoding="utf-8")
            print(json.dumps({"output": str(args.out.resolve()), "status": result["status"], "evidence_count": len(result["evidence"])}))
        else:
            print(content, end="")
        return 2 if result["status"] == "failed" else 0
    except Exception as exc:
        print(json.dumps({"schema_version": 2, "status": "failed", "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
